"""PPT 兜底渲染器（v2.0 大点 02 降级为"快速模式"，不再是默认链路）。

什么时候用：
- `POST /api/ppt/{id}/generate`（快速模式：无 Agent、无配图、无自检，秒级出片）；
- OfficeCLI 不可用（`PPT_OFFICECLI_PATH` 配错/缺失）时的最后兜底（验证项 7）。

与 Agent 链路的关系（裁定 8）：**同一任务只会走其中一条**。Agent 链路的唯一写者是
OfficeCLI；本模块是 python-pptx 渲染器，两者不并发写同一个文件。

修掉的老问题（裁定 11）：
- 要点不再 `points[:5]` 静默截断：>5 条自动双栏，>10 条才截断（并 warning）；
- 目录页按容量分两栏，章节多时不再溢出页外；
- 版面尺寸统一 16:9，内容过多时按条数降字号（20/18/16）。

对外接口（保持不变，供 ppt_service 调用）：
    generate_local_ppt(outline, topic, out_path, template_path=None) -> int  # 返回页数
"""
import logging
from pathlib import Path

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

logger = logging.getLogger(__name__)

NAVY = RGBColor(0x0E, 0x2A, 0x47)
NAVY_LIGHT = RGBColor(0x1F, 0x38, 0x64)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0x59, 0x59, 0x59)
FONT = "微软雅黑"
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)

MAX_BULLETS = 10     # 单页要点上限（超过才截断）
TOC_PER_COLUMN = 5   # 目录每栏最多 5 个章节


# --------------------------------------------------------------------------- #
# 基础绘制
# --------------------------------------------------------------------------- #
def _add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _add_rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _add_text(slide, x, y, w, h, text, size=18, color=NAVY_LIGHT, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _add_bullets(slide, x, y, w, h, bullets, size):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, pt in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(max(6, size // 2))
        run = p.add_run()
        run.text = f"▸ {pt}"
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = NAVY_LIGHT
    return box


def _header(slide, title):
    _add_rect(slide, 0, 0, SLIDE_W, Inches(1.0), NAVY)
    _add_rect(slide, 0, Inches(1.0), SLIDE_W, Inches(0.06), GOLD)
    _add_text(slide, Inches(0.6), Inches(0.16), Inches(11.6), Inches(0.7),
              title[:40], size=26, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def _page_no(slide, num):
    _add_text(slide, Inches(12.0), Inches(7.0), Inches(0.9), Inches(0.35),
              str(num), size=12, color=GRAY, align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------- #
# 各类页面
# --------------------------------------------------------------------------- #
def _render_cover(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, NAVY)
    _add_rect(slide, 0, Inches(6.95), SLIDE_W, Inches(0.55), GOLD)
    _add_rect(slide, Inches(0.9), Inches(2.05), Inches(0.09), Inches(1.6), GOLD)
    _add_text(slide, Inches(1.15), Inches(2.0), Inches(11), Inches(1.5),
              title[:30], size=44, color=WHITE, bold=True)
    _add_text(slide, Inches(1.15), Inches(3.65), Inches(11), Inches(0.6),
              f"—— {subtitle} ——", size=20, color=GOLD)
    _add_text(slide, Inches(1.15), Inches(6.25), Inches(8.5), Inches(0.45),
              "蜀道集团安全管理部 ｜ 安全培训课件", size=14, color=WHITE)


def _render_toc(prs, sections, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _header(slide, "目 录  CONTENTS")
    for i, sec in enumerate(sections[: TOC_PER_COLUMN * 2]):
        col, row = divmod(i, TOC_PER_COLUMN)
        x = 1.2 + col * 5.9
        y = 1.55 + row * 1.02
        _add_rect(slide, Inches(x), Inches(y), Inches(0.72), Inches(0.72),
                  NAVY if i % 2 == 0 else GOLD)
        _add_text(slide, Inches(x), Inches(y), Inches(0.72), Inches(0.72),
                  f"{i + 1:02d}", size=22, color=WHITE, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _add_text(slide, Inches(x + 0.95), Inches(y), Inches(4.3), Inches(0.72),
                  str(sec.get("name") or f"章节 {i + 1}")[:22], size=20,
                  color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    _page_no(slide, num)


def _render_section(prs, name, idx, num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, NAVY)
    _add_rect(slide, 0, Inches(6.95), SLIDE_W, Inches(0.55), GOLD)
    _add_text(slide, Inches(1.0), Inches(2.35), Inches(3.0), Inches(1.6),
              f"{idx:02d}", size=96, color=GOLD, bold=True)
    _add_text(slide, Inches(1.0), Inches(4.05), Inches(11), Inches(0.9),
              str(name)[:26], size=36, color=WHITE, bold=True)


def _render_content(prs, title, bullets, num):
    """内容页：>5 条要点走双栏；按条数/长度选字号（不再静默截断）。"""
    bullets = [str(b).strip() for b in bullets if str(b).strip()]
    if len(bullets) > MAX_BULLETS:
        logger.warning("单页要点 %d 条超过上限 %d，截断（标题：%s）", len(bullets), MAX_BULLETS, title)
        bullets = bullets[:MAX_BULLETS]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _header(slide, title)
    _add_rect(slide, Inches(0.62), Inches(1.45), Inches(0.08), Inches(5.2), GOLD)

    longest = max((len(b) for b in bullets), default=0)
    if len(bullets) > 5 or longest > 34:
        size = 18 if len(bullets) <= 8 else 16
        half = (len(bullets) + 1) // 2
        _add_bullets(slide, Inches(0.95), Inches(1.55), Inches(5.6), Inches(5.2),
                     bullets[:half], size)
        _add_bullets(slide, Inches(6.85), Inches(1.55), Inches(5.6), Inches(5.2),
                     bullets[half:], size)
    else:
        size = 20 if len(bullets) <= 4 else 18
        _add_bullets(slide, Inches(0.95), Inches(1.55), Inches(11.3), Inches(5.2),
                     bullets, size)
    _page_no(slide, num)


def _render_ending(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, NAVY)
    _add_rect(slide, 0, Inches(6.95), SLIDE_W, Inches(0.55), GOLD)
    _add_text(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.1),
              "高高兴兴上班来，平平安安回家去", size=38, color=WHITE, bold=True,
              align=PP_ALIGN.CENTER)
    _add_text(slide, Inches(1.0), Inches(3.85), Inches(11.3), Inches(0.6),
              f"《{title}》· 感谢聆听", size=20, color=GOLD, align=PP_ALIGN.CENTER)
    _add_text(slide, Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.45),
              "蜀道集团安全管理部 ｜ 安全培训课件", size=13, color=WHITE,
              align=PP_ALIGN.CENTER)


# --------------------------------------------------------------------------- #
# 入口
# --------------------------------------------------------------------------- #
def generate_local_ppt(
    outline: dict, topic: str, out_path: Path, template_path: Path | None = None
) -> int:
    """渲染大纲 → pptx，返回页数。

    outline 结构：{"title", "subtitle", "sections": [{"name", "points"}]}
    （也被 Agent 的 outline.json 复用，但 Agent 走 OfficeCLI，不调用这里）

    template_path：用户上传的 .pptx。**保留其首页/母版/样式**，在末尾追加页面
    （历史行为，快速模式沿用；Agent 链路不追加而是只借其主题色）。
    """
    title = (outline.get("title") or topic).strip()
    subtitle = (outline.get("subtitle") or "安全培训讲座课件").strip()
    sections = [s for s in (outline.get("sections") or []) if s.get("name") or s.get("points")]

    if template_path is not None and Path(template_path).exists():
        prs = Presentation(str(template_path))
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    _render_cover(prs, title, subtitle)
    _render_toc(prs, sections, 2)
    num = 3
    for i, sec in enumerate(sections, start=1):
        _render_section(prs, sec.get("name") or f"章节 {i}", i, num)
        num += 1
        points = sec.get("points") or []
        if points:
            _render_content(prs, sec.get("name") or f"章节 {i}", points, num)
            num += 1
    _render_ending(prs, title)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    page_count = len(prs.slides._sldIdLst)
    logger.info(
        "本地兜底 PPT 生成完成: %s（%d 页，模板=%s）", out_path, page_count, template_path or "内置简洁"
    )
    return page_count

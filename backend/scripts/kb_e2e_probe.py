"""04 真机 E2E：上传 docx（含表格）→ 管理员审核 → 落 knowledge_src → 检索命中。

不连 HTTP：直接调用 service（等价于路由内部逻辑），但落盘、转换、清洗、查重、入库全真实。
"""
import sys
import uuid
from pathlib import Path

B = Path(r"D:\code\aicode\ai智能体开发实战5期\weekend\ShuDao Smart AI Security Assistant\team\backend")
sys.path.insert(0, str(B))
import os
os.chdir(B)

from src.ai.knowledge import upsert_chunks  # noqa: E402
from src.core.config import settings  # noqa: E402
from src.core.db import SessionLocal  # noqa: E402
from src.core.security import hash_password  # noqa: E402
from src.models.user import User  # noqa: E402
from src.schemas.knowledge_contrib import ReviewRequest  # noqa: E402
from src.services.knowledge_contrib_service import KnowledgeContribService  # noqa: E402


class FakeUpload:
    def __init__(self, path: Path, filename: str):
        self.filename = filename
        self.file = open(path, "rb")


def main() -> None:
    tag = uuid.uuid4().hex[:6]
    title = f"塔吊安装拆卸安全操作规程（04验证{tag}）"
    body = ("第{}条 塔吊安装前应检查基础承载力与预埋件，附墙装置每 6m 设置一道，"
            "顶升作业应有专人指挥并设置警戒区，风速超过 12m/s 禁止作业。\n")

    # 1) 造一份含表格的 docx（模拟真实上传件）
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    for i in range(1, 30):
        doc.add_paragraph(body.format(i))
    table = doc.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("项目", "要求"), ("附墙间距", "每 6m 一道"), ("顶升风速", "≤12m/s")]):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    src = Path(settings.UPLOAD_DIR) / f"kb_contrib/e2e_{tag}.docx"
    src.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(src))
    print(f"[1] 生成 docx: {src.name}（{src.stat().st_size} 字节）")

    db = SessionLocal()
    try:
        # 2) 造一个测试上传者
        user = User(username=f"e2e_kc_{tag}", hashed_password=hash_password("pass123456"), role="employee")
        db.add(user)
        db.commit()
        db.refresh(user)
        admin_id = db.query(User.id).filter(User.role == "admin").first()[0]
        print(f"[2] 测试用户 #{user.id} / 管理员 #{admin_id}")

        # 3) 提交
        contrib = KnowledgeContribService.submit(
            db, title, "procedures", "04 端到端验证：塔吊安装规程，含表格。", FakeUpload(src, src.name), user.id
        )
        print(f"[3] 提交成功 贡献 #{contrib.id}，状态 {contrib.status}")

        # 4) 审核通过
        out = KnowledgeContribService.review(
            db, contrib.id, ReviewRequest(approved=True, note="04 E2E"), {"id": admin_id, "role": "admin"}
        )
        print(f"[4] 审核通过 → 状态 {out.status}，块数 {out.chunk_count}，转换 {out.convert_by}")
        print(f"    落盘路径: {out.md_path}")

        md_file = Path(settings.KNOWLEDGE_SRC_DIR) / Path(out.md_path).relative_to("knowledge_src")
        print(f"[5] 文件存在: {md_file.is_file()}，大小 {md_file.stat().st_size if md_file.exists() else 0} 字节")
        if md_file.is_file():
            head = md_file.read_text(encoding="utf-8").splitlines()[:6]
            print("    前 6 行:", head)
            print("    含表格:", any(line.startswith("|") for line in head + md_file.read_text(encoding='utf-8').splitlines()[:80]))

        # 5) 检索命中（真嵌入 + Chroma）
        from src.ai.rag import retrieve_queries

        blocks = retrieve_queries(["塔吊附墙间距要求"])
        hit = [b for b in blocks if b.get("metadata", {}).get("doc") == md_file.stem]
        print(f"[6] 检索 4 条中命中本贡献: {len(hit)} 条")
        if hit:
            md = hit[0]["metadata"]
            print("    metadata:", {k: md.get(k) for k in ("doc", "origin", "contrib_id", "doc_type", "source")})
            print("    片段:", (hit[0].get("text") or "")[:80].replace("\n", " "))

        # 6) 查重：同标题再提交应被拒
        try:
            dup = KnowledgeContribService.submit(
                db, title, "procedures", "重复标题测试", FakeUpload(src, src.name), user.id
            )
            KnowledgeContribService.review(
                db, dup.id, ReviewRequest(approved=True, note="dup"), {"id": admin_id, "role": "admin"}
            )
            print("[7] 查重：未拦截（异常）")
        except Exception as exc:  # noqa: BLE001
            print(f"[7] 查重拦截: {type(exc).__name__}: {str(exc)[:120]}")
    finally:
        db.close()


if __name__ == "__main__":
    main()

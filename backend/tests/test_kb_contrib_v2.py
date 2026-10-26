"""v2.0 大点 04：kb_convert（格式转换）+ contrib_store（落盘/查重/回滚）单测。

这些用例**不加载嵌入模型、不连 Chroma、不调 MinerU**：
- 转换走"内置兜底"路径（MinerU 未安装时的真实分支）；
- 入库用假 `upsert_file`，只验证文件写入、查重与回滚顺序。
"""
from pathlib import Path
from types import SimpleNamespace

import pytest

from src.services import contrib_store, kb_convert


# --------------------------------------------------------------------------- #
# 转换：md / txt / html / docx / pdf
# --------------------------------------------------------------------------- #
def test_md_direct_strips_front_matter(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("---\ntitle: 旧标题\n---\n\n# 已有标题\n\n正文内容", encoding="utf-8")
    r = kb_convert.to_markdown(f)
    assert r["converted_by"] == "direct"
    assert r["md"].startswith("# 已有标题")
    assert "title: 旧标题" not in r["md"]


def test_txt_gets_heading(tmp_path):
    f = tmp_path / "高处作业要点.txt"
    f.write_text("第一条 为了加强高处作业安全管理……", encoding="utf-8")
    r = kb_convert.to_markdown(f)
    assert r["converted_by"] == "builtin"
    assert r["md"].splitlines()[0] == "# 高处作业要点"


def test_html_converts_table_and_drops_script(tmp_path):
    f = tmp_path / "b.html"
    f.write_text(
        "<html><head><style>.x{}</style></head><body><h1>高处作业措施</h1>"
        "<table><tr><th>项目</th><th>要求</th></tr><tr><td>安全带</td><td>高挂低用</td></tr></table>"
        "<script>var a=1</script></body></html>",
        encoding="utf-8",
    )
    r = kb_convert.to_markdown(f)
    assert "| 项目 | 要求 |" in r["md"]
    assert "var a" not in r["md"]
    assert r["md"].startswith("# 高处作业措施")


def _make_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("塔吊安装规程", level=1)
    doc.add_paragraph("第一条 安装前应检查基础。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "附墙"
    table.cell(1, 1).text = "每 6m 一道"
    doc.save(str(path))


def test_docx_fallback_keeps_heading_and_table(tmp_path, monkeypatch):
    monkeypatch.setattr(kb_convert.settings, "KB_MINERU_ENABLED", False)
    f = tmp_path / "塔吊安装规程.docx"
    _make_docx(f)
    r = kb_convert.to_markdown(f)
    assert r["converted_by"] == "fallback"
    assert "# 塔吊安装规程" in r["md"]
    assert "| 项目 | 要求 |" in r["md"]
    assert "每 6m 一道" in r["md"]
    assert r["warnings"]  # 提示表格可能不完整


def test_pdf_fallback_extracts_text(tmp_path, monkeypatch):
    import pymupdf

    monkeypatch.setattr(kb_convert.settings, "KB_MINERU_ENABLED", False)
    f = tmp_path / "高处作业.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 80), "High-risk work above 2 meters requires a harness.", fontsize=12)
    page.insert_text((60, 110), "第二条 临边应设置防护栏杆。", fontsize=12)
    doc.save(str(f))
    doc.close()

    r = kb_convert.to_markdown(f)
    assert r["converted_by"] == "fallback"
    assert "harness" in r["md"]


def test_scanned_pdf_is_rejected(tmp_path, monkeypatch):
    import pymupdf

    monkeypatch.setattr(kb_convert.settings, "KB_MINERU_ENABLED", False)
    f = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    doc.new_page()  # 空白页 = 无文本层
    doc.save(str(f))
    doc.close()
    with pytest.raises(ValueError, match="无可复制文本|为空"):
        kb_convert.to_markdown(f)


def test_unsupported_ext(tmp_path):
    f = tmp_path / "x.docx"
    f.write_bytes(b"PK")
    f2 = tmp_path / "x.xlsx"
    f2.write_bytes(b"PK")
    with pytest.raises(ValueError, match="不支持"):
        kb_convert.to_markdown(f2)


# --------------------------------------------------------------------------- #
# 1000 字门槛（裁定 3）
# --------------------------------------------------------------------------- #
def test_quality_min_chars_configurable():
    from src.services import kb_govern

    short = "第1条 施工现场安全要求。" * 32  # 约 416 字：过 300、不到 1000
    assert kb_govern.quality_check(short)["ok"] is True          # 历史 300 门槛
    qc = kb_govern.quality_check(short, min_chars=1000)
    assert qc["ok"] is False and any("碎片阈值 1000" in r for r in qc["reasons"])


# --------------------------------------------------------------------------- #
# 落盘 / 查重 / 回滚
# --------------------------------------------------------------------------- #
@pytest.fixture()
def kb_dir(tmp_path, monkeypatch):
    """把 knowledge_src 指到临时目录，避免污染真实知识库。"""
    root = tmp_path / "knowledge_src"
    (root / "laws").mkdir(parents=True)
    monkeypatch.setattr(contrib_store.settings, "KNOWLEDGE_SRC_DIR", str(root))
    return root


class _FakeDB:
    """只实现 contrib_store 用到的 scalars()（查 DB 已入库标题）。"""

    def __init__(self, rows=None):
        self._rows = rows or []

    def scalars(self, _stmt):
        rows = self._rows

        class _R:
            def all(self_inner):
                return rows

        return _R()

    def add(self, _obj):
        return None


def _contrib(**kw):
    base = dict(
        id=99,
        title="塔吊安装安全规程",
        doc_type="laws",
        origin_filename="塔吊安装安全规程.docx",
        md_path=None,
        chunk_count=None,
        convert_by=None,
        source_kind="user",
        source_agent_id=None,
        origin_urls=None,
    )
    base.update(kw)
    return SimpleNamespace(**base)


def test_normalize_stem_strips_ext_and_bad_chars():
    assert contrib_store.normalize_stem("塔吊/安装:规程?.docx") == "安装 规程"
    assert contrib_store.normalize_stem("") == "未命名文档"
    assert len(contrib_store.normalize_stem("长" * 200)) == 60


def test_store_contribution_writes_and_ingests(kb_dir, monkeypatch):
    calls = {}

    def fake_upsert_file(path, doc_type, *, stem=None, source=None, extra_meta=None):
        calls.update(path=str(path), doc_type=doc_type, stem=stem, source=source, meta=extra_meta)
        assert Path(path).name.startswith(".tmp_")      # 先写临时文件
        return 3

    monkeypatch.setattr(contrib_store, "upsert_file", fake_upsert_file)
    db = _FakeDB()
    c = _contrib()
    out = contrib_store.store_contribution(db, c, "# 塔吊安装安全规程\n\n正文")

    assert out["chunks"] == 3 and out["stem"] == "塔吊安装安全规程"
    final = kb_dir / "laws" / "塔吊安装安全规程.md"
    assert final.is_file()
    assert not list((kb_dir / "laws").glob(".tmp_*"))
    assert c.md_path == "knowledge_src/laws/塔吊安装安全规程.md"
    assert calls["doc_type"] == "laws" and calls["meta"]["contrib_id"] == 99
    assert calls["source"].startswith("knowledge_src/laws/")


def test_conflict_same_filename(kb_dir, monkeypatch):
    (kb_dir / "laws" / "重复标题.md").write_text("# 重复标题\n\n旧内容", encoding="utf-8")
    monkeypatch.setattr(contrib_store, "upsert_file", lambda *a, **k: 1)
    with pytest.raises(contrib_store.ConflictError) as ei:
        contrib_store.store_contribution(
            _FakeDB(),
            _contrib(title="重复标题", origin_filename="重复标题.md"),
            "# 重复标题\n\n新内容",
        )
    assert ei.value.kind == "file"


def test_conflict_same_title_other_file(kb_dir, monkeypatch):
    (kb_dir / "laws" / "existing.md").write_text("# 塔吊安装安全规程\n\n官方原文", encoding="utf-8")
    monkeypatch.setattr(contrib_store, "upsert_file", lambda *a, **k: 1)
    with pytest.raises(contrib_store.ConflictError) as ei:
        contrib_store.store_contribution(
            _FakeDB(),
            _contrib(origin_filename="另一个文件名.md"),
            "# 塔吊安装安全规程\n\n正文",
        )
    assert ei.value.kind == "title"


def test_rollback_deletes_tmp_and_chunks(kb_dir, monkeypatch):
    removed = {}

    def boom(path, doc_type, **kw):
        raise RuntimeError("嵌入模型炸了")

    monkeypatch.setattr(contrib_store, "upsert_file", boom)
    monkeypatch.setattr(contrib_store, "remove_chunks", lambda doc: removed.setdefault("doc", doc))
    with pytest.raises(RuntimeError):
        contrib_store.store_contribution(
            _FakeDB(), _contrib(), "# 塔吊安装安全规程" + chr(10) * 2 + "正文"
        )
    assert not list((kb_dir / "laws").glob(".tmp_*"))
    assert not (kb_dir / "laws" / "塔吊安装安全规程.md").exists()
